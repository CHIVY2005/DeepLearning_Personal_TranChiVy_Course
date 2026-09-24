from pathlib import Path
from tempfile import TemporaryDirectory
import math
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "BT01_Regression_TranChiVy_3123580065.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in (
        ("Title", 24, "17365D"),
        ("Heading 1", 16, "17365D"),
        ("Heading 2", 14, "1F4E79"),
        ("Heading 3", 12, "2F5597"),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    formula = doc.styles.add_style("Formula", WD_STYLE_TYPE.PARAGRAPH)
    formula.font.name = "Cambria Math"
    formula._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    formula.font.size = Pt(11.5)
    formula.paragraph_format.left_indent = Cm(0.7)
    formula.paragraph_format.right_indent = Cm(0.4)
    formula.paragraph_format.space_before = Pt(3)
    formula.paragraph_format.space_after = Pt(3)

    result = doc.styles.add_style("Result", WD_STYLE_TYPE.PARAGRAPH)
    result.font.name = "Times New Roman"
    result._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    result.font.size = Pt(12)
    result.font.bold = True
    result.font.color.rgb = RGBColor.from_string("006100")
    result.paragraph_format.left_indent = Cm(0.7)
    result.paragraph_format.space_before = Pt(4)
    result.paragraph_format.space_after = Pt(7)

    for sec in doc.sections:
        add_page_number(sec.footer.paragraphs[0])
    return doc


def add_centered(doc, text, size=12, bold=False, color=None, after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_formula(doc, text):
    return doc.add_paragraph(text, style="Formula")


def add_result(doc, text):
    return doc.add_paragraph("Kết quả: " + text, style="Result")


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for j, value in enumerate(headers):
        cell = header.cells[j]
        cell.text = str(value)
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(10.5)
    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = str(value)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10.5)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_question(doc, number, title, prompt=None):
    doc.add_heading(f"Câu {number}. {title}", level=2)
    if prompt:
        p = doc.add_paragraph()
        r = p.add_run("Đề bài tóm tắt: ")
        r.bold = True
        p.add_run(prompt)


def add_cover(doc):
    add_centered(doc, "BÀI TẬP 01", 18, True, "17365D", 4)
    add_centered(doc, "REGRESSION", 26, True, "1F4E79", 18)
    add_centered(doc, "Hồi quy tuyến tính – Hồi quy phân lớp", 15, True, "404040", 36)

    box = doc.add_table(rows=4, cols=2)
    box.style = "Table Grid"
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    info = [
        ("Họ và tên sinh viên", "Trần Chí Vỹ"),
        ("Mã số sinh viên", "3123580065"),
        ("Học phần", "Deep Learning"),
        ("Nội dung", "Ví dụ trong slide và 02 bộ bài tập Regression"),
    ]
    for i, (label, value) in enumerate(info):
        box.cell(i, 0).text = label
        box.cell(i, 1).text = value
        set_cell_shading(box.cell(i, 0), "D9EAF7")
        for run in box.cell(i, 0).paragraphs[0].runs:
            run.bold = True
        for cell in box.rows[i].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)
    doc.add_paragraph()
    add_centered(doc, "Tháng 09 năm 2026", 12, False, "666666", 0)
    doc.add_page_break()


def add_intro(doc):
    doc.add_heading("GHI CHÚ TRÌNH BÀY", level=1)
    add_bullets(doc, [
        "Ký hiệu ŷ là giá trị dự đoán; e = y − ŷ là phần dư, trừ khi đề quy định khác.",
        "Logarithm trong cross-entropy và log-odds là logarithm tự nhiên ln.",
        "Các kết quả thập phân được làm tròn 4–6 chữ số; phép tính trung gian giữ đủ độ chính xác.",
        "Với nghiệm bình phương tối thiểu, ma trận thiết kế có cột 1 để biểu diễn hệ số chặn.",
    ])
    doc.add_heading("BỐ CỤC BÀI LÀM", level=1)
    add_bullets(doc, [
        "Phần I: Ví dụ mẫu trong slide – giải hệ phương trình chuẩn, pseudo-inverse và Gradient Descent.",
        "Phần II: Bài tập 1 – 10 câu trong DL_01_Regression.docx.",
        "Phần III: Bài tập 2 – 15 câu trong DL_01_Regression_P2.docx.",
    ])
    doc.add_page_break()


def add_slide_example(doc):
    doc.add_heading("PHẦN I. VÍ DỤ MẪU TRONG SLIDE", level=1)
    doc.add_paragraph(
        "Dữ liệu chiều cao x (cm) và cân nặng y (kg) của 5 người trong slide. "
        "Tìm mô hình ŷ = ax + b làm cực tiểu tổng bình phương phần dư."
    )
    data = [("A", 170, 65), ("B", 155, 50), ("C", 150, 45), ("D", 175, 70), ("E", 165, 55)]
    add_table(doc, ["Mẫu", "x", "y", "x²", "xy"], [
        (name, x, y, x*x, x*y) for name, x, y in data
    ] + [("Tổng", 815, 285, 133275, 46875)])

    doc.add_heading("1. Giải hệ phương trình chuẩn", level=2)
    add_formula(doc, "R(a,b) = Σᵢ(axᵢ + b − yᵢ)²")
    doc.add_paragraph("Cho ∂R/∂a = 0 và ∂R/∂b = 0, ta được:")
    add_formula(doc, "133275a + 815b = 46875")
    add_formula(doc, "815a + 5b = 285")
    doc.add_paragraph(
        "Từ phương trình thứ hai: b = 57 − 163a. Thế vào phương trình thứ nhất:"
    )
    add_formula(doc, "133275a + 815(57 − 163a) = 46875 ⇒ 430a = 420")
    add_formula(doc, "a = 42/43 = 0,976744;  b = 57 − 163(42/43) = −4395/43 = −102,209302")
    add_result(doc, "ŷ = 0,976744x − 102,209302.")

    doc.add_heading("2. Giải bằng pseudo-inverse", level=2)
    doc.add_paragraph("Đặt Y = Aθ với θ = [a, b]ᵀ:")
    add_formula(doc, "A = [[170,1],[155,1],[150,1],[175,1],[165,1]],  Y = [65,50,45,70,55]ᵀ")
    add_formula(doc, "AᵀA = [[133275,815],[815,5]],  AᵀY = [46875,285]ᵀ")
    add_formula(doc, "det(AᵀA) = 133275·5 − 815² = 2150 ≠ 0")
    add_formula(doc, "(AᵀA)⁻¹ = (1/2150)[[5,−815],[−815,133275]]")
    add_formula(doc, "A⁺ = (AᵀA)⁻¹Aᵀ")
    add_table(doc, ["Hàng của A⁺", "A", "B", "C", "D", "E"], [
        ("Hàng 1", "0,016279", "−0,018605", "−0,030233", "0,027907", "0,004651"),
        ("Hàng 2", "−2,453488", "3,232558", "5,127907", "−4,348837", "−0,558140"),
    ])
    add_formula(doc, "θ = A⁺Y = (AᵀA)⁻¹AᵀY = [0,976744; −102,209302]")
    add_result(doc, "Pseudo-inverse cho đúng cùng nghiệm: a = 0,976744; b = −102,209302.")

    doc.add_heading("3. Giải bằng Gradient Descent", level=2)
    doc.add_paragraph(
        "Do x có độ lớn khoảng 150–175, chuẩn hóa x giúp chọn learning rate ổn định. "
        "Đặt z = (x − 163)/√86; khi đó mean(z)=0 và mean(z²)=1. Viết ŷ = αz + c."
    )
    add_formula(doc, "J(α,c) = (1/2m)Σ(αzᵢ + c − yᵢ)²,  m = 5")
    add_formula(doc, "∂J/∂α = (1/m)Σ(ŷᵢ−yᵢ)zᵢ;  ∂J/∂c = (1/m)Σ(ŷᵢ−yᵢ)")
    doc.add_paragraph(
        "Khởi tạo α₀=c₀=0, η=0,1. Vì z đã được chuẩn hóa, α* = 9,057953 và c* = 57; "
        "mỗi bước có dạng αₜ₊₁=0,9αₜ+0,1α*, cₜ₊₁=0,9cₜ+5,7."
    )
    sigma = math.sqrt(86)
    alpha_star = (42/43) * sigma
    gd_rows = []
    for t in (0, 1, 2, 5, 10, 50, 100):
        f = 1 - 0.9**t
        alpha = alpha_star * f
        c = 57 * f
        a = alpha / sigma
        b = c - 163*a
        gd_rows.append((t, f"{alpha:.6f}", f"{c:.6f}", f"{a:.6f}", f"{b:.6f}"))
    add_table(doc, ["Vòng lặp t", "αₜ", "cₜ", "aₜ=αₜ/√86", "bₜ=cₜ−163aₜ"], gd_rows)
    doc.add_paragraph(
        "Khi t→∞: α→9,057953 và c→57. Đổi về biến gốc: a=α/√86=0,976744; "
        "b=c−163a=−102,209302."
    )
    add_result(doc, "Gradient Descent hội tụ đến ŷ = 0,976744x − 102,209302, trùng với hai cách trên.")
    doc.add_page_break()


def add_exercise_set_1(doc, chart_path):
    doc.add_heading("PHẦN II. BÀI TẬP 1", level=1)

    add_question(doc, 1, "Xây dựng phương trình hồi quy tuyến tính đơn")
    add_table(doc, ["x", "y", "x−x̄", "y−ȳ", "Tích", "(x−x̄)²"], [
        (1,3,-2,-4,8,4),(2,5,-1,-2,2,1),(3,7,0,0,0,0),(4,9,1,2,2,1),(5,11,2,4,8,4),
        ("Σx=15","Σy=35","","",20,10)
    ])
    add_formula(doc, "x̄ = 15/5 = 3;  ȳ = 35/5 = 7")
    add_formula(doc, "SSxy = 20; SSx = 10; β₁ = SSxy/SSx = 2; β₀ = ȳ − β₁x̄ = 1")
    add_result(doc, "ŷ = 1 + 2x; khi x=6 thì ŷ=13.")

    add_question(doc, 2, "Tính Mean Squared Error (MSE)")
    add_table(doc, ["x", "y", "ŷ=2x+1", "e=y−ŷ", "e²"], [
        (1,4,3,1,1),(2,6,5,1,1),(3,8,7,1,1),(4,11,9,2,4)
    ])
    add_formula(doc, "MSE = (1+1+1+4)/4 = 7/4 = 1,75")
    add_result(doc, "Sai số trung bình bình phương khá nhỏ, nhưng mô hình dự đoán thấp hơn thực tế ở cả 4 mẫu (có độ lệch âm của dự đoán).")

    add_question(doc, 3, "Gradient Descent cho Linear Regression")
    doc.add_paragraph("Tại w=0, b=0, các dự đoán đều bằng 0.")
    add_table(doc, ["x", "y", "ŷ", "ŷ−y", "(ŷ−y)x"], [(1,3,0,-3,-3),(2,5,0,-5,-10),(4,9,0,-9,-36)])
    add_formula(doc, "∂J/∂w = (1/3)Σ(ŷ−y)x = −49/3 = −16,333333")
    add_formula(doc, "∂J/∂b = (1/3)Σ(ŷ−y) = −17/3 = −5,666667")
    add_formula(doc, "w₁ = 0 − 0,1(−49/3) = 1,633333; b₁ = 0 − 0,1(−17/3) = 0,566667")
    add_result(doc, "Sau 1 bước: ŷ = 1,633333x + 0,566667.")

    add_question(doc, 4, "Hồi quy tuyến tính đa biến")
    add_formula(doc, "X = [[1,1,2],[1,2,1],[1,3,2],[1,4,3]],  Y = [5,6,9,12]ᵀ")
    add_formula(doc, "XᵀX = [[4,10,8],[10,30,22],[8,22,18]]")
    add_formula(doc, "XᵀY = [32,92,70]ᵀ")
    add_formula(doc, "β = (XᵀX)⁻¹XᵀY = [1,2,1]ᵀ")
    add_result(doc, "ŷ = 1 + 2x₁ + x₂; tại (x₁,x₂)=(5,2), ŷ=13.")

    add_question(doc, 5, "Hồi quy tuyến tính – một bước Batch Gradient Descent")
    add_formula(doc, "J(0,0) = (1/10)(2²+3²+5²+4²+6²) = 9")
    add_formula(doc, "∂J/∂w = −(1/5)Σxᵢyᵢ = −69/5 = −13,8")
    add_formula(doc, "∂J/∂b = −(1/5)Σyᵢ = −20/5 = −4")
    add_formula(doc, "w₁ = 0−0,1(−13,8)=1,38; b₁=0−0,1(−4)=0,4")
    add_result(doc, "Mô hình sau bước 1: ŷ=1,38x+0,4; với x=6, ŷ=8,68.")

    add_question(doc, 6, "Phân lớp bằng đường quyết định tuyến tính")
    add_table(doc, ["Điểm", "(x₁,x₂)", "2x₁−x₂+1", "Lớp"], [
        ("A","(1,1)",2,1),("B","(2,1)",4,1),("C","(1,3)",0,1),("D","(4,2)",7,1)
    ])
    add_formula(doc, "Đường quyết định: 2x₁ − x₂ + 1 = 0 ⇔ x₂ = 2x₁ + 1")
    if chart_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(chart_path), width=Inches(5.8))
        cap = doc.add_paragraph("Hình 1. Các điểm dữ liệu và đường phân lớp.")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
    add_result(doc, "A, B, C, D đều thuộc lớp 1; điểm C nằm đúng trên biên và được xếp lớp 1 theo quy tắc z≥0.")

    add_question(doc, 7, "Logistic Regression – tính xác suất phân lớp")
    add_formula(doc, "z = 0,5·2 − 1·1 + 0,2 = 0,2")
    add_formula(doc, "P(y=1|x) = σ(0,2) = 1/(1+e⁻⁰·²) = 0,549834")
    add_result(doc, "Ngưỡng 0,5: lớp 1. Ngưỡng 0,8: lớp 0.")

    add_question(doc, 8, "Tính Cross-Entropy Loss")
    add_table(doc, ["Mẫu", "y", "p", "Loss"], [
        (1,1,0.9,"−ln(0,9)=0,105361"),(2,0,0.2,"−ln(0,8)=0,223144"),(3,1,0.7,"−ln(0,7)=0,356675")
    ])
    add_formula(doc, "Tổng loss = 0,685179; average loss = 0,685179/3 = 0,228393")
    add_result(doc, "Average loss thấp; các xác suất nhìn chung phù hợp nhãn, trong đó mẫu 3 kém chắc chắn nhất.")

    add_question(doc, 9, "Perceptron Learning")
    doc.add_paragraph(
        "Dùng quy ước sign(z)=+1 khi z≥0 (nhất quán với Câu 10). Ban đầu score=0 nên hai mẫu dương đúng, "
        "mẫu (1,3), t=−1 bị phân lớp sai."
    )
    add_table(doc, ["Mẫu", "score trước cập nhật", "Dự đoán", "Xử lý", "w, b sau mẫu"], [
        ("(1,1), t=+1",0,"+1","Đúng","(0,0), 0"),
        ("(2,1), t=+1",0,"+1","Đúng","(0,0), 0"),
        ("(1,3), t=−1",0,"+1","Sai: w←w+tx; b←b+t","(−1,−3), −1"),
    ])
    add_formula(doc, "Đường quyết định mới: −x₁ − 3x₂ − 1 = 0")
    doc.add_paragraph(
        "Lưu ý: một số giáo trình coi score=0 là chưa phân loại và cập nhật khi t·score≤0. Theo quy ước đó, "
        "mẫu 1 cũng được cập nhật; sau 1 epoch nhận w=(0,−2), b=0."
    )
    add_result(doc, "Theo quy tắc sign(z)=+1 tại z=0: w=(−1,−3), b=−1.")

    add_question(doc, 10, "Hồi quy phân lớp – một bước Perceptron theo gradient")
    doc.add_paragraph("Hiểu ký hiệu w₀=𝟙 trong đề là vector w₀=(1,1), b₀=0; learning rate ρ=1.")
    add_table(doc, ["Mẫu", "t", "score=wᵀx+b", "Dự đoán", "Đúng/Sai"], [
        ("(1,1)",1,2,1,"Đúng"),("(2,1)",1,3,1,"Đúng"),("(1,2)",-1,3,1,"Sai"),("(2,2)",-1,4,1,"Sai")
    ])
    add_formula(doc, "Y = {((1,2),−1), ((2,2),−1)}")
    add_formula(doc, "J(w,b)=−Σₖ∈Y tₖ(wᵀxₖ+b)")
    add_formula(doc, "∂J/∂w = −Σₖ∈Y tₖxₖ = (3,4);  ∂J/∂b = −Σₖ∈Y tₖ = 2")
    add_formula(doc, "w₁ = w₀ − ρ∂J/∂w = (1,1)−(3,4)=(−2,−3); b₁=0−2=−2")
    add_result(doc, "Sau một bước batch Perceptron: w₁=(−2,−3), b₁=−2; biên: −2x₁−3x₂−2=0.")
    doc.add_page_break()


def add_exercise_set_2(doc):
    doc.add_heading("PHẦN III. BÀI TẬP 2", level=1)

    add_question(doc, 1, "Tính hệ số hồi quy tuyến tính đơn bằng OLS")
    add_table(doc, ["x", "y", "x²", "xy"], [(1,1.2,1,1.2),(2,1.8,4,3.6),(3,2.6,9,7.8),(4,3.2,16,12.8),(5,3.8,25,19.0),("Σx=15","Σy=12,6","Σx²=55","Σxy=44,4")])
    add_formula(doc, "x̄=3; ȳ=2,52; a₁=(44,4−5·3·2,52)/(55−5·3²)=6,6/10=0,66")
    add_formula(doc, "a₀=ȳ−a₁x̄=2,52−0,66·3=0,54")
    add_result(doc, "ŷ=0,54+0,66x; tuần 7: ŷ=5,16; tuần 9: ŷ=6,48.")

    add_question(doc, 2, "Tính sai số dự đoán của mô hình hồi quy")
    add_table(doc, ["x", "y", "ŷ=0,54+0,66x", "e=y−ŷ", "e²"], [
        (1,1.2,1.20,0.00,0.0000),(2,1.8,1.86,-0.06,0.0036),(3,2.6,2.52,0.08,0.0064),(4,3.2,3.18,0.02,0.0004),(5,3.8,3.84,-0.04,0.0016)
    ])
    add_formula(doc, "SSE = 0+0,0036+0,0064+0,0004+0,0016 = 0,012")
    add_result(doc, "SSE rất nhỏ nên mô hình khớp tốt với 5 điểm dữ liệu.")

    add_question(doc, 3, "Tính MAE, MSE, RMSE")
    add_table(doc, ["Mẫu", "y", "ŷ", "|e|", "e²"], [(1,80,75,5,25),(2,75,85,10,100),(3,90,88,2,4),(4,100,95,5,25)])
    add_formula(doc, "MAE=(5+10+2+5)/4=5,5")
    add_formula(doc, "MSE=(25+100+4+25)/4=38,5; RMSE=√38,5=6,204837")
    add_result(doc, "Sai số điển hình khoảng 6,20 đơn vị; mẫu 2 có sai số lớn nhất (10 đơn vị).")

    add_question(doc, 4, "Tính Relative MSE và Coefficient of Variation")
    add_formula(doc, "ȳ_train=(80+90+100+110+120)/5=100")
    add_formula(doc, "SSE_test=(80−75)²+(75−85)²=125")
    add_formula(doc, "Σ(yᵢ−ȳ_train)²=(80−100)²+(75−100)²=1025")
    add_formula(doc, "RelMSE=125/1025=0,121951")
    add_formula(doc, "RMSE=√(125/2)=7,905694; CV=7,905694/100=0,079057≈7,91%")
    add_result(doc, "RelMSE≈0,122<1 và CV≈7,91% cho thấy sai số tương đối thấp so với mức doanh số trung bình.")

    add_question(doc, 5, "Hồi quy tuyến tính dạng ma trận")
    add_formula(doc, "X=[[1,1],[1,2],[1,3],[1,4]],  Y=[1,3,4,8]ᵀ")
    add_formula(doc, "XᵀX=[[4,10],[10,30]]")
    add_formula(doc, "(XᵀX)⁻¹=(1/20)[[30,−10],[−10,4]] = [ [1,5  −0,5]; [−0,5  0,2] ]")
    add_formula(doc, "XᵀY=[16,51]ᵀ; a=(XᵀX)⁻¹XᵀY=[−1,5; 2,2]ᵀ")
    add_result(doc, "ŷ = −1,5 + 2,2x.")

    add_question(doc, 6, "Hồi quy tuyến tính đa biến")
    add_formula(doc, "X=[[1,1,4],[1,2,5],[1,3,8],[1,4,2]],  Y=[1,6,8,12]ᵀ")
    add_formula(doc, "XᵀX=[[4,10,19],[10,30,46],[19,46,109]]; XᵀY=[27,85,122]ᵀ")
    add_formula(doc, "a=(XᵀX)⁻¹XᵀY=[−311/183, 425/122, −10/183]ᵀ")
    add_formula(doc, "a≈[−1,699454; 3,483607; −0,054645]")
    add_result(doc, "ŷ=−1,699454+3,483607x₁−0,054645x₂; tại (5,6), ŷ=15,390710.")

    add_question(doc, 7, "Hồi quy đa thức bậc 2")
    add_table(doc, ["x", "y", "x²", "x³", "x⁴", "xy", "x²y"], [
        (1,1,1,1,1,1,1),(2,4,4,8,16,8,16),(3,9,9,27,81,27,81),(4,15,16,64,256,60,240),("Σx=10","Σy=29","Σx²=30","Σx³=100","Σx⁴=354","Σxy=96","Σx²y=338")
    ])
    add_formula(doc, "Hệ chuẩn: [[4,10,30],[10,30,100],[30,100,354]][a₀,a₁,a₂]ᵀ=[29,96,338]ᵀ")
    add_formula(doc, "Giải hệ được a₀=−3/4=−0,75; a₁=19/20=0,95; a₂=3/4=0,75")
    add_result(doc, "ŷ=−0,75+0,95x+0,75x²; khi x=5, ŷ=22,75.")

    add_question(doc, 8, "Logistic Regression – tính xác suất và phân lớp")
    rows=[]
    for x in (40,60,80):
        z=-4+0.08*x
        p=1/(1+math.exp(-z))
        rows.append((x,f"{z:.1f}",f"{p:.6f}","Đậu" if p>=0.5 else "Rớt"))
    add_table(doc,["Điểm x","z=−4+0,08x","p=σ(z)","Ngưỡng 0,5"],rows)
    add_result(doc, "x=40: Rớt; x=60 và x=80: Đậu.")

    add_question(doc, 9, "Tính odds và log-odds")
    add_table(doc,["Email","p","odds=p/(1−p)","logit(p)=ln(odds)"],[
        ("A",0.2,0.25,"−1,386294"),("B",0.5,1,"0"),("C",0.8,4,"1,386294")
    ])
    add_result(doc, "Khi p tăng từ 0 đến 1, odds và log-odds đều tăng; log-odds âm khi p<0,5, bằng 0 tại p=0,5 và dương khi p>0,5.")

    add_question(doc, 10, "Logistic Regression với hai biến đầu vào")
    students=[("A",60,50),("B",70,80),("C",40,45),("D",90,85)]
    rows=[]
    for name,x1,x2 in students:
        z=-3+0.04*x1+0.06*x2
        p=1/(1+math.exp(-z))
        rows.append((name,x1,x2,f"{z:.1f}",f"{p:.6f}","Trúng tuyển" if p>=0.5 else "Không","Nhận" if p>=0.8 else "Không nhận"))
    add_table(doc,["SV","x₁","x₂","z","p","Ngưỡng 0,5","Ngưỡng 0,8"],rows)
    add_result(doc, "Theo ngưỡng 0,5: cả A, B, C, D đều trúng tuyển. Nếu yêu cầu p≥0,8: nhận A, B và D; không nhận C.")

    add_question(doc, 11, "Tính hệ số xác định R²")
    add_formula(doc, "ȳ=(10+15+20+25+30)/5=20")
    add_formula(doc, "SST=(10−20)²+(15−20)²+(20−20)²+(25−20)²+(30−20)²=250")
    add_formula(doc, "SSE=(10−12)²+(15−14)²+(20−19)²+(25−27)²+(30−28)²=14")
    add_formula(doc, "R²=1−SSE/SST=1−14/250=0,944")
    add_result(doc, "Mô hình giải thích khoảng 94,4% biến thiên của y; mức độ phù hợp cao.")

    add_question(doc, 12, "Ridge Regression – hàm mất mát có phạt")
    add_formula(doc, "Penalty=λa₁²=2·3²=18; Loss=40+18=58")
    add_formula(doc, "Khi λ=5: penalty=5·9=45; Loss=40+45=85")
    add_result(doc, "Tăng λ làm tăng mức phạt hệ số lớn, co nhỏ hệ số và giảm phương sai/overfitting, nhưng λ quá lớn có thể gây underfitting.")

    add_question(doc, 13, "LASSO Regression – hàm mất mát có phạt")
    add_formula(doc, "Σ|aᵢ|=|3|+|−2|+|0,5|=5,5")
    add_formula(doc, "Penalty=λΣ|aᵢ|=2·5,5=11; Loss=40+11=51")
    add_result(doc, "Phạt L1 tạo nghiệm thưa và có thể đưa một số hệ số đúng về 0; do đó LASSO thực hiện chọn đặc trưng.")

    add_question(doc, 14, "Elastic Net – loss kết hợp Ridge và LASSO")
    add_formula(doc, "Thành phần LASSO: λ₁Σ|aᵢ|=1·(2+1+3)=6")
    add_formula(doc, "Thành phần Ridge: λ₂Σaᵢ²=0,5·(4+1+9)=7")
    add_formula(doc, "Loss=50+6+7=63")
    add_result(doc, "Elastic Net phù hợp khi vừa cần chọn đặc trưng vừa cần ổn định hệ số, đặc biệt khi các biến đầu vào tương quan cao.")

    add_question(doc, 15, "Bài tập tổng hợp")
    add_table(doc,["Tuần","x","y","ŷ=10+5x","|e|","e²"],[(1,2,20,20,0,0),(2,3,25,25,0,0),(3,5,35,35,0,0),(4,7,45,45,0,0),(5,9,55,55,0,0)])
    add_formula(doc, "x̄=5,2; ȳ=36; SSxy=Σ(x−x̄)(y−ȳ)=164; SSx=32,8")
    add_formula(doc, "a₁=164/32,8=5; a₀=36−5·5,2=10 ⇒ ŷ=10+5x")
    add_formula(doc, "Tại x=10: ŷ=60")
    add_formula(doc, "MAE=0; MSE=0; RMSE=0; SSE=0; R²=1−0/SST=1")
    add_result(doc, "Chi phí quảng cáo và doanh số có quan hệ tuyến tính dương hoàn hảo trong tập dữ liệu: tăng 1 đơn vị quảng cáo gắn với tăng 5 đơn vị doanh số.")

    doc.add_paragraph()
    end = doc.add_paragraph("— HẾT —")
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    end.runs[0].bold = True


def create_chart(path):
    os.environ.setdefault("MPLCONFIGDIR", str(path.parent / "mpl-cache"))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    pts = {"A":(1,1), "B":(2,1), "C":(1,3), "D":(4,2)}
    fig, ax = plt.subplots(figsize=(6.2,4.2), dpi=180)
    xs = [-0.5, 4.5]
    ax.plot(xs, [2*x+1 for x in xs], color="#C00000", linewidth=2, label="2x₁ − x₂ + 1 = 0")
    for name,(x1,x2) in pts.items():
        ax.scatter([x1],[x2], s=55, color="#1F4E79")
        ax.annotate(name, (x1,x2), xytext=(5,5), textcoords="offset points", fontsize=10)
    ax.set_xlim(-0.5,4.5); ax.set_ylim(0,10.5)
    ax.set_xlabel("x₁"); ax.set_ylabel("x₂")
    ax.grid(True, alpha=0.25); ax.legend(loc="upper left")
    fig.tight_layout(); fig.savefig(path, bbox_inches="tight"); plt.close(fig)


def main():
    doc = setup_document()
    add_cover(doc)
    add_intro(doc)
    with TemporaryDirectory() as tmp:
        chart = Path(tmp) / "decision_boundary.png"
        create_chart(chart)
        add_slide_example(doc)
        add_exercise_set_1(doc, chart)
        add_exercise_set_2(doc)
        doc.core_properties.title = "BT01 Regression – Trần Chí Vỹ – 3123580065"
        doc.core_properties.subject = "Linear Regression and Linear Classification"
        doc.core_properties.author = "Trần Chí Vỹ"
        doc.core_properties.keywords = "Regression, OLS, Pseudo-inverse, Gradient Descent, Logistic Regression"
        doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
